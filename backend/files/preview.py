"""表格预览（设计文档 9.3）：CSV / Excel 只取前若干行，单元格一律转成文字，交给前端画表格。"""

import codecs
import csv
import io
from datetime import date, datetime, time
from pathlib import PurePosixPath
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel

from common.errors import InvalidRequestError

PREVIEW_ROWS = 100
MAX_COLUMNS = 50
MAX_PREVIEW_BYTES = 20 * 1024 * 1024  # 更大的文件整读进内存不划算，请用户下载查看
_DELIMITERS = {".csv": ",", ".tsv": "\t"}
_EXCEL = frozenset({".xlsx", ".xlsm"})
# csv 模块默认单个单元格最多 128KB，长文本列会直接报错；预览的文件本身不超过 20MB
csv.field_size_limit(MAX_PREVIEW_BYTES)


class TablePreview(BaseModel):
    sheets: list[str]  # Excel 的工作表名（按顺序）；CSV 为空
    sheet: str | None  # 当前显示的工作表
    rows: list[list[str]]  # 前若干行（第一行通常是表头），每行补齐到同样的列数
    truncated: bool  # 后面还有没显示的行
    columns_truncated: bool = False  # 有一行用到的列超过 MAX_COLUMNS：只显示前面这些列


class PreviewError(InvalidRequestError):
    """不能按表格预览；消息直接给用户看（统一映射成 422，积压小问题 I2）。"""


def table_preview(
    name: str, data: bytes, *, sheet: str | None = None, limit: int = PREVIEW_ROWS
) -> TablePreview:
    suffix = PurePosixPath(name).suffix.lower()
    if suffix in _DELIMITERS:
        utf16 = data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE))
        # 文字文件里不会有空字节：多半是二进制文件改了后缀。带文件头标记的 UTF-16（Excel"Unicode
        # 文本"一类）每个英文字符都带一个空字节，按 UTF-16 读（积压小问题 D9）
        if not utf16 and b"\x00" in data[:4096]:
            raise PreviewError("不是有效的 CSV 文件")
        # newline=""：交给 csv 模块自己识别换行（含只用回车换行的旧 Mac 文件）
        text = io.StringIO(
            data.decode("utf-16", errors="replace") if utf16 else _decode(data), newline=""
        )
        reader = csv.reader(text, delimiter=_DELIMITERS[suffix])
        try:
            rows = [row for _, row in zip(range(limit + 1), reader, strict=False)]
        except csv.Error as exc:  # 如二进制内容改了后缀
            raise PreviewError("不是有效的 CSV 文件") from exc
        return _preview(rows, limit, sheets=[], sheet=None)
    if suffix in _EXCEL:
        return _excel(data, sheet, limit)
    raise PreviewError("这种文件不能按表格预览")


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")  # Excel 导出的 UTF-8 文件常带 BOM
    except UnicodeDecodeError:
        return data.decode("gb18030", errors="replace")  # 中文系统上的 Excel 导出常是 GBK


def _excel(data: bytes, sheet: str | None, limit: int) -> TablePreview:
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, OSError) as exc:
        raise PreviewError("不是有效的 Excel 文件") from exc
    try:
        names = book.sheetnames
        name = sheet or names[0]
        if name not in names:
            raise PreviewError(f"工作表不存在：{name}")
        values: list[tuple[object, ...]] = list(
            book[name].iter_rows(max_row=limit + 1, values_only=True)
        )
    finally:
        book.close()
    if any(value is None for row in values for value in row):
        values = _with_formulas(data, name, limit, values)
    rows = [[_text(value) for value in row] for row in values]
    return _preview(rows, limit, sheets=names, sheet=name)


def _with_formulas(
    data: bytes, sheet: str, limit: int, values: list[tuple[object, ...]]
) -> list[tuple[object, ...]]:
    """没有缓存结果的公式格（程序写的 xlsx 常这样）显示公式原文，不是空格子（积压小问题 D8）。
    算出结果要另加公式计算的依赖，这里不做；只读前 limit 行，多读一遍代价小。"""
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    try:
        formulas = list(book[sheet].iter_rows(max_row=limit + 1, values_only=True))
    finally:
        book.close()
    return [
        tuple(
            formula
            if value is None and isinstance(formula, str) and formula.startswith("=")
            else value
            for value, formula in zip(row, formula_row, strict=False)
        )
        for row, formula_row in zip(values, formulas, strict=False)
    ]


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date | time):
        return value.isoformat()
    return str(value)


def _preview(
    rows: list[list[str]], limit: int, *, sheets: list[str], sheet: str | None
) -> TablePreview:
    shown = rows[:limit]
    used = max((_used_width(row) for row in shown), default=0)
    width = min(MAX_COLUMNS, used)
    return TablePreview(
        sheets=sheets,
        sheet=sheet,
        rows=[(row + [""] * width)[:width] for row in shown],
        truncated=len(rows) > limit,
        columns_truncated=used > MAX_COLUMNS,
    )


def _used_width(row: list[str]) -> int:
    """去掉行尾空单元格后的列数（Excel 的整行常带一长串空格子）。"""
    for index in range(len(row), 0, -1):
        if row[index - 1]:
            return index
    return 0
