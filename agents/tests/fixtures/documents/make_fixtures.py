"""生成文档解析测试用的样例文件（二进制文件随测试保留，脚本只在需要重新生成时运行）。

pymupdf、python-docx 不是项目依赖，用临时环境运行：
uv run --no-project --with pymupdf==1.28.2 --with python-docx==1.2.0 --with python-pptx==1.0.2 \\
    --with openpyxl==3.1.5 python agents/tests/fixtures/documents/make_fixtures.py
"""

from pathlib import Path

import docx
import openpyxl
import pymupdf
from pptx import Presentation

HERE = Path(__file__).resolve().parent


def pdf() -> None:
    """3 页，每页一个章节标题；正文带精确词 ERR_ORDER_n00。"""
    document = pymupdf.open()
    for n in range(1, 4):
        page = document.new_page()
        page.insert_text((72, 72), f"第{n}章 订单系统设计", fontname="china-s", fontsize=18)
        body = (
            f"这是第 {n} 页的正文。订单服务负责下单、支付回调与退款，"
            f"错误码 ERR_ORDER_{n}00 表示库存不足。"
        )
        page.insert_textbox(
            pymupdf.Rect(72, 100, 520, 700), body * 3, fontname="china-s", fontsize=11
        )
    document.save(HERE / "design.pdf")


def scanned_pdf() -> None:
    """只有图形、没有文字（相当于扫描件）。"""
    document = pymupdf.open()
    page = document.new_page()
    page.draw_rect(pymupdf.Rect(50, 50, 200, 200), fill=(0.5, 0.5, 0.5))
    document.save(HERE / "scan.pdf")


def word() -> None:
    document = docx.Document()
    document.add_heading("员工手册", 0)
    document.add_heading("第一章 考勤", 1)
    document.add_paragraph("上班时间为 9:30 至 18:30，弹性 30 分钟。")
    document.add_heading("1.1 请假", 2)
    document.add_paragraph("病假需提供医院证明，年假每年 15 天。")
    document.add_heading("第二章 报销", 1)
    document.add_paragraph("差旅报销需在 30 天内提交。")
    document.save(HERE / "handbook.docx")


def slides() -> None:
    deck = Presentation()
    for title, body in [
        ("季度复盘", "营收增长 20%\n新增客户 300 家"),
        ("下季度计划", "上线知识库\n优化检索"),
    ]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    deck.save(str(HERE / "review.pptx"))


def sheets() -> None:
    book = openpyxl.Workbook()
    sales = book.worksheets[0]
    sales.title = "销售"
    for row in (["月份", "销售额"], ["一月", 120], ["二月", 98]):
        sales.append(row)
    book.create_sheet("成本").append(["项目", "金额"])
    book.save(HERE / "sales.xlsx")


def others() -> None:
    (HERE / "faq.html").write_text(
        "<html><head><title>FAQ</title><script>alert(1)</script></head><body>"
        "<h1>常见问题</h1><h2>如何退款</h2><p>在订单页点击退款。</p></body></html>",
        encoding="utf-8",
    )
    # 老格式 Word：只有 OLE 文件头，markitdown 不认识
    (HERE / "old.doc").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600)


if __name__ == "__main__":
    pdf()
    scanned_pdf()
    word()
    slides()
    sheets()
    others()
